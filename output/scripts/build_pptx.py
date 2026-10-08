"""Build HeteroCam Lab Chinese presentation with native PPT shapes + diagrams."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsmap, qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
FIG = Path(__file__).resolve().parents[1] / "assets" / "figures"
OUT = Path(__file__).resolve().parents[1] / "final_presentation_cn.pptx"

W, H = Inches(13.333), Inches(7.5)
BG = RGBColor(0x0E, 0x14, 0x1C)
INK = RGBColor(0xE8, 0xEE, 0xF4)
MUTED = RGBColor(0x8B, 0x9A, 0xA8)
GOLD = RGBColor(0xE4, 0xC9, 0x8A)
CYAN = RGBColor(0x6E, 0xE0, 0xC5)
ORANGE = RGBColor(0xF0, 0xA3, 0x6B)
BLUE = RGBColor(0x7E, 0xB6, 0xFF)
CARD = RGBColor(0x1A, 0x23, 0x2E)
LINE = RGBColor(0x2A, 0x35, 0x40)
FONT = "微软雅黑"

L = Inches(0.55)
TITLE_T = Inches(0.28)
TITLE_H = Inches(0.62)


def _set_run(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rpr, qn("a:ea"))
    ea.set("typeface", FONT)


def _fill_slide(slide, color=BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _tb(slide, l, t, w, h, text, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size, color, bold)
    return box


def _title(slide, text):
    return _tb(slide, L, TITLE_T, Inches(12.2), TITLE_H, text, size=26, color=GOLD, bold=True)


def _notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def _rect(slide, l, t, w, h, fill, line=None, line_w=Pt(1.25)):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = line or fill
    sh.line.width = line_w
    sh.adjustments[0] = 0.08
    return sh


def _shape_text(sh, text, size=12, color=INK, bold=False, align=PP_ALIGN.CENTER):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.alignment = align
    # clear default
    if p.runs:
        p.runs[0].text = ""
    run = p.add_run()
    run.text = text
    _set_run(run, size, color, bold)
    try:
        tf._txBody.bodyPr.set("anchor", "ctr")
    except Exception:
        pass


def _pic(slide, name, l, t, w, h):
    p = FIG / name
    if p.exists():
        slide.shapes.add_picture(str(p), l, t, w, h)
        return True
    return False


def _footer(slide, page, n=16):
    _tb(slide, L, Inches(7.18), Inches(8), Inches(0.24), "HeteroCam Lab  ·  异构摄像集群交互仿真", 10, MUTED)
    _tb(slide, Inches(11.4), Inches(7.18), Inches(1.4), Inches(0.24), f"{page} / {n}", 10, MUTED, align=PP_ALIGN.RIGHT)


def new_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_slide(s)
    return s


def slide_title(prs):
    s = new_slide(prs)
    _tb(s, L, Inches(1.7), Inches(12), Inches(0.4), "HETEROCAM LAB", 14, CYAN, True)
    _tb(s, L, Inches(2.15), Inches(12.2), Inches(1.3), "异构摄像集群交互仿真", 40, INK, True)
    _tb(
        s,
        L,
        Inches(3.55),
        Inches(12),
        Inches(0.7),
        "具身运动约束  ·  针孔多视几何  ·  可遥控指挥台",
        18,
        MUTED,
    )
    chips = [
        ("PTZ 固定云台", GOLD),
        ("RAIL 滑轨", RGBColor(0xC5, 0xCD, 0xD8)),
        ("ROVER 拍摄机器人", ORANGE),
        ("DOG 机器狗", CYAN),
        ("UAV 无人机", BLUE),
    ]
    for i, (lab, c) in enumerate(chips):
        sh = _rect(s, L + Inches(i * 2.45), Inches(4.7), Inches(2.28), Inches(0.7), CARD, c)
        _shape_text(sh, lab, 12, c, True)
    _tb(s, L, Inches(5.7), Inches(12), Inches(0.4), "浏览器运动学仿真  ·  Vite / React / Three.js / zustand", 13, MUTED)
    _notes(s, "开场说明：这是一套可在浏览器里交互的异构摄像集群仿真，不是动力学物理引擎，目标是拍摄任务规划、覆盖与内外参这一层。")
    _footer(s, 1)


def slide_problem(prs):
    s = new_slide(prs)
    _title(s, "拍摄集群不能当成「会飞的相同相机」")
    items = [
        ("具身", "PTZ 不能走，轮式不能横移，狗脖子行程有限，无人机才有高度。控制律必须按构型空间分开。", GOLD),
        ("视觉", "每台设备是针孔相机：内参随变焦变，外参由底盘+云台推出，还要处理遮挡和多视基线。", CYAN),
        ("交互", "能选中、遥控、改任务模式、看各机实时画面，并把 K 与 T_wc 导出给后续 CV 管线。", ORANGE),
    ]
    for i, (h, body, c) in enumerate(items):
        sh = _rect(s, L, Inches(1.15 + i * 1.7), Inches(12.2), Inches(1.52), CARD, c)
        _tb(s, L + Inches(0.35), Inches(1.28 + i * 1.7), Inches(1.4), Inches(1.2), h, 22, c, True, anchor=MSO_ANCHOR.MIDDLE)
        _tb(s, L + Inches(2.0), Inches(1.35 + i * 1.7), Inches(9.8), Inches(1.15), body, 16, INK, anchor=MSO_ANCHOR.MIDDLE)
    _notes(s, "强调三个约束同时成立。若只做相同无人机编队，或只做单相机预览，都覆盖不了本题。")
    _footer(s, 2)


def slide_system(prs):
    s = new_slide(prs)
    _title(s, "系统：指挥台驱动仿真核，五类平台协同拍同一主体")
    _pic(s, "fig_system.png", Inches(0.35), Inches(0.95), Inches(12.6), Inches(5.95))
    _notes(s, "从上到下是人、HUD、zustand 核、五类平台、摄影棚主体。颜色与软件界面中的机位配色一致。")
    _footer(s, 3)


def slide_why(prs):
    s = new_slide(prs)
    _title(s, "做成浏览器运动学仿真，是为了当场能改机位")
    headers = ["方案", "代价", "对本题"]
    rows = [
        ["ROS + Gazebo / Isaac", "URDF、物理、安装链很长", "交互周期不适合空仓库交付"],
        ["Unity", "编辑器授权与资源管线", "难在本仓库直接打开"],
        ["浏览器 Three.js", "无额外运行时", "运动学、相机、UI 同一套代码"],
    ]
    col_w = [Inches(3.4), Inches(4.3), Inches(4.4)]
    x0, y0, rh = L, Inches(1.2), Inches(1.05)
    for j, h in enumerate(headers):
        x = x0 + sum(col_w[:j], Inches(0))
        sh = _rect(s, x, y0, col_w[j] - Inches(0.08), Inches(0.55), CARD, GOLD)
        _shape_text(sh, h, 13, GOLD, True)
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            x = x0 + sum(col_w[:j], Inches(0))
            ec = CYAN if i == 2 else LINE
            sh = _rect(s, x, y0 + Inches(0.7 + i * 1.05), col_w[j] - Inches(0.08), rh - Inches(0.1), CARD, ec)
            _shape_text(sh, cell, 14, INK if i == 2 else MUTED, i == 2)
    _tb(s, L, Inches(5.7), Inches(12.2), Inches(0.9), "碰撞是圆柱 vs AABB，不求解接触力。对覆盖规划、跟踪模式和内外参导出已经够用。", 15, MUTED)
    _notes(s, "选型理由要讲清楚：不是否定物理仿真，而是本题的决策层在运动学与多视几何。")
    _footer(s, 4)


def slide_framework(prs):
    s = new_slide(prs)
    _title(s, "五层框架：几何先于网格，60 Hz 姿态不进 React")
    _pic(s, "fig_framework.png", Inches(0.4), Inches(0.95), Inches(12.5), Inches(6.0))
    _notes(s, "最底层是类型和世界常量，往上才是相机几何、tick、R3F、HUD。平台位姿用 useFrame 写 ref，只有增删机位才改 agentIds。")
    _footer(s, 5)


def slide_modules(prs):
    s = new_slide(prs)
    _title(s, "模块切开后，可见性可以不启动 WebGL 就推理")
    _pic(s, "fig_modules.png", Inches(0.4), Inches(0.95), Inches(12.5), Inches(6.0))
    _notes(s, "左列 sim/* 不依赖 Three。右列只负责画出来和接收输入。箭头表示渲染读仿真状态。")
    _footer(s, 6)


def slide_loop(prs):
    s = new_slide(prs)
    _title(s, "一帧里先步进，再离屏拍各机，最后才画指挥视角")
    _pic(s, "fig_loop.png", Inches(0.35), Inches(1.05), Inches(12.6), Inches(5.7))
    _notes(s, "优先级 −3 的时钟保证几何先于 FBO。dt 钳在 0.05，避免切后台后一帧走出棚。")
    _footer(s, 7)


def slide_kin(prs):
    s = new_slide(prs)
    _title(s, "五类平台的构型空间不同，控制律必须分开写")
    _pic(s, "fig_kinematics.png", Inches(0.3), Inches(0.95), Inches(12.7), Inches(6.05))
    _notes(s, "PTZ 无位移；滑轨一维；ROVER 是 unicycle 先转后走；狗可侧移但脖子行程小；无人机才有高度。默认 DOG-1 为手动，开机就能 WASD。")
    _footer(s, 8)


def slide_camera(prs):
    s = new_slide(prs)
    _title(s, "光学中心由底盘加云台推出，变焦只缩小垂直视场")
    _pic(s, "fig_camera.png", Inches(0.25), Inches(0.95), Inches(12.8), Inches(5.95))
    _notes(s, "Three.js 中 rotation.x>0 是抬头。实现时曾把 tilt 取反，滑轨和无人机全部朝天。导出的 T_wc 不是 OpenCV 的 T_cw。")
    _footer(s, 9)


def slide_coverage(prs):
    s = new_slide(prs)
    _title(s, "LOCK 要过视锥和遮挡两关，覆盖热力只采样视锥")
    _pic(s, "fig_coverage.png", Inches(0.3), Inches(1.05), Inches(12.7), Inches(5.7))
    _notes(s, "热力约 4 Hz 更新且不算遮挡，是为了交互还能满帧。三角化分数是相对 60° 的启发量，不要当成协方差。")
    _footer(s, 10)


def slide_render(prs):
    s = new_slide(prs)
    _title(s, "三个图层把调试图元挡在机体镜头之外")
    _pic(s, "fig_render.png", Inches(0.3), Inches(1.0), Inches(12.7), Inches(5.8))
    _notes(s, "第一版用 drei Hud 的正交相机 makeDefault，主视角被换成正交相机后近乎全黑。现把画面条作为轨道相机的子物体，并放在 layer 2。")
    _footer(s, 11)


def slide_interact(prs):
    s = new_slide(prs)
    _title(s, "三种任务模式共用同一套云台，但谁有权改角度不同")
    modes = [
        ("手动", "键盘与滑条有效。双击地面给路点，移动平台驶向该点。选中时不自动 look-at。", GOLD),
        ("跟踪", "每帧把云台对准主体胸部。PTZ / 滑轨主要靠这一模式干活。", CYAN),
        ("巡航", "ROVER / DOG / UAV 绕主体；滑轨跟 x；云台仍跟踪。", ORANGE),
    ]
    for i, (h, body, c) in enumerate(modes):
        sh = _rect(s, L + Inches(i * 4.15), Inches(1.15), Inches(3.95), Inches(2.55), CARD, c)
        _tb(s, L + Inches(i * 4.15) + Inches(0.22), Inches(1.3), Inches(3.5), Inches(0.5), h, 20, c, True)
        _tb(s, L + Inches(i * 4.15) + Inches(0.22), Inches(1.85), Inches(3.5), Inches(1.6), body, 14, INK)
    keys = [
        ("WASD / QE", "机动与升降"),
        ("方向键 · ZX", "云台与变焦"),
        ("空格 F G B", "暂停与图层"),
        ("1–8 单击", "选机 / 画面条"),
    ]
    for i, (k, v) in enumerate(keys):
        sh = _rect(s, L + Inches(i * 3.1), Inches(4.0), Inches(2.95), Inches(1.35), CARD, LINE)
        _tb(s, L + Inches(i * 3.1) + Inches(0.18), Inches(4.12), Inches(2.6), Inches(0.45), k, 14, GOLD, True)
        _tb(s, L + Inches(i * 3.1) + Inches(0.18), Inches(4.55), Inches(2.6), Inches(0.6), v, 13, MUTED)
    _tb(s, L, Inches(5.55), Inches(12.2), Inches(1.1), "环绕编队只重排 ROVER / DOG / UAV，PTZ 与滑轨保持原位。导出 JSON 含 K、T_wc、是否看见主体、GSD。", 15, MUTED)
    _notes(s, "演示时可先跟踪看 LOCK，再切手动用 WASD 推狗，最后点环绕编队。")
    _footer(s, 12)


def slide_path(prs):
    s = new_slide(prs)
    _title(s, "实现顺序：先能算看见，再往视口上挂网格")
    steps = [
        ("1", "Vite 空项目"),
        ("2", "types → kinematics\n→ coverage → store"),
        ("3", "棚、主体、五类外形"),
        ("4", "FBO 多路画面"),
        ("5", "HUD 与导出"),
        ("6", "浏览器里修\n朝向 / 雾 / 图层"),
    ]
    for i, (n, lab) in enumerate(steps):
        x = L + Inches(i * 2.1)
        c = _rect(s, x, Inches(1.35), Inches(1.9), Inches(2.1), CARD, GOLD if i in (1, 5) else CYAN)
        _tb(s, x, Inches(1.45), Inches(1.9), Inches(0.45), n, 20, GOLD, True, PP_ALIGN.CENTER)
        _tb(s, x + Inches(0.08), Inches(1.95), Inches(1.74), Inches(1.3), lab, 13, INK, align=PP_ALIGN.CENTER)
        if i < 5:
            _tb(s, x + Inches(1.78), Inches(2.1), Inches(0.35), Inches(0.4), "→", 16, MUTED)
    bugs = [
        ("主视角近黑", "Hud 正交相机 makeDefault 抢走默认相机 → 画面条改挂轨道相机 + layer 2"),
        ("高机位朝天", "tilt 取反，与 Three.js 抬头符号相反 → tilt = 仰角"),
        ("滑轨拍幕墙", "轨道在 −Z，默认 yaw=0 也看 −Z → 滑轨 yaw=π"),
        ("铭牌把 Canvas 掀翻", "drei Html 在 React 19 卸载 portal → Billboard + Text"),
    ]
    for i, (h, b) in enumerate(bugs):
        y = Inches(3.7) + Inches(i * 0.78)
        _tb(s, L, y, Inches(2.4), Inches(0.7), h, 14, ORANGE, True, anchor=MSO_ANCHOR.MIDDLE)
        _tb(s, L + Inches(2.5), y, Inches(9.7), Inches(0.7), b, 14, INK, anchor=MSO_ANCHOR.MIDDLE)
    _notes(s, "排障四条都是真实交付阻塞，不是装饰性「经验总结」。汇报时可以对照现场演示。")
    _footer(s, 13)


def slide_ui(prs):
    s = new_slide(prs)
    _title(s, "指挥台把编队、指标和各机针孔画面放在同一屏")
    ui = FIG / "ui_capture.png"
    if ui.exists():
        s.shapes.add_picture(str(ui), Inches(0.45), Inches(1.0), Inches(12.4), Inches(5.55))
    else:
        zones = [
            (0.55, 1.15, 2.8, 5.2, "左栏编队", "LOCK / LOST\n加入 PTZ…UAV"),
            (3.5, 1.15, 6.3, 3.55, "三维视口", "主体 8 字 · 视锥 · 热力 · 基线"),
            (9.95, 1.15, 2.8, 5.2, "右栏检视", "云台滑条 · K/fy\n覆盖率 · GSD"),
            (3.5, 4.85, 6.3, 1.5, "底部六路针孔画面", "点击可选中对应机位"),
        ]
        for x, y, w, h, t, sub in zones:
            sh = _rect(s, Inches(x), Inches(y), Inches(w), Inches(h), CARD, CYAN if "视口" in t else GOLD)
            _tb(s, Inches(x) + Inches(0.2), Inches(y) + Inches(0.2), Inches(w - 0.3), Inches(0.4), t, 16, GOLD, True)
            _tb(s, Inches(x) + Inches(0.2), Inches(y) + Inches(0.7), Inches(w - 0.3), Inches(h - 1.0), sub, 14, MUTED)
        _tb(s, L, Inches(6.55), Inches(12), Inches(0.35), "运行 npm run dev 打开 localhost:5173，可对照本页结构操作。", 12, MUTED)
    _notes(s, "若现场能投屏仿真，这一页可以很快带过，把时间留给 WASD 和环绕编队。")
    _footer(s, 14)


def slide_export(prs):
    s = new_slide(prs)
    _title(s, "导出的是世界系光学中心，不是 OpenCV 的相机系")
    left = _rect(s, L, Inches(1.15), Inches(6.1), Inches(5.15), CARD, GOLD)
    _tb(s, L + Inches(0.3), Inches(1.35), Inches(5.5), Inches(0.4), "JSON 里有什么", 16, GOLD, True)
    _tb(
        s,
        L + Inches(0.3),
        Inches(1.85),
        Inches(5.5),
        Inches(4.1),
        "time / subject\n每机 name, kind, mode\nT_wc.t  光学中心\nrpy_yxz  (pitch, yaw, 0)\nK：fx=fy, cx, cy, 分辨率, fovY\nseesSubject、gsd_cm_per_px",
        16,
        INK,
    )
    right = _rect(s, Inches(6.95), Inches(1.15), Inches(5.8), Inches(5.15), CARD, ORANGE)
    _tb(s, Inches(7.25), Inches(1.35), Inches(5.3), Inches(0.4), "刻意没做的", 16, ORANGE, True)
    _tb(
        s,
        Inches(7.25),
        Inches(1.85),
        Inches(5.3),
        Inches(4.1),
        "无 IMU / 轮速噪声\n无径向畸变与卷帘\n热力不算遮挡\n离屏机位线性增加 GPU 负担\n上限 8 台，避免画面条挤爆",
        16,
        INK,
    )
    _notes(s, "给做多视几何的同事时，第一句话先讲坐标系。需要 PnP 就在脚本里把 T_wc 求逆。")
    _footer(s, 15)


def slide_end(prs):
    s = new_slide(prs)
    _title(s, "能改机位、能看见、能把位姿拿走")
    pts = [
        "五类异构平台共用一个针孔相机模型和同一主体。",
        "仿真核与渲染分离：LOCK 是几何算出来的，不是画出来的。",
        "指挥台把遥控、覆盖指标和六路画面放在同一决策环里。",
        "可续：T_cw 矩阵、深度附件、可配置巡航半径。",
    ]
    for i, t in enumerate(pts):
        _tb(s, L, Inches(1.25 + i * 1.05), Inches(12.2), Inches(0.9), t, 20, INK, anchor=MSO_ANCHOR.MIDDLE)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, L, Inches(1.25 + i * 1.05), Inches(0.08), Inches(0.85))
        bar.fill.solid()
        bar.fill.fore_color.rgb = [GOLD, CYAN, ORANGE, BLUE][i]
        bar.line.fill.background()
    _tb(s, L, Inches(5.6), Inches(12), Inches(0.5), "文档：docs/heterocam-lab-technical.md    运行：npm run dev", 14, MUTED)
    _notes(s, "收束到三个可演示事实：改机位、LOCK 判定、JSON 导出。后续工作只点深度和坐标形式，不要展开成新项目。")
    _footer(s, 16)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide_title(prs)
    slide_problem(prs)
    slide_system(prs)
    slide_why(prs)
    slide_framework(prs)
    slide_modules(prs)
    slide_loop(prs)
    slide_kin(prs)
    slide_camera(prs)
    slide_coverage(prs)
    slide_render(prs)
    slide_interact(prs)
    slide_path(prs)
    slide_ui(prs)
    slide_export(prs)
    slide_end(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print("saved", OUT)


if __name__ == "__main__":
    main()
